"""A spreadsheet read as Excel lays it out: what the canvas draws for an Excel file.

Each cell keeps what Excel stores (a number, a date, text) with its number format, so the web app
shows `25%` for 0.25 and `1.234,50` for 1234.5 in the owner's language, as Excel does. The layout
comes with it: merged cells, column widths, row heights, hidden rows and columns, and each cell's
font, fill, alignment and borders. An invoice laid out on a grid reads like the invoice.

The file is from outside and parsing it is synchronous: callers run ``read_workbook`` in a worker
thread, and anything that goes wrong inside a file is ``Unreadable``, never a crash. `.xlsx` cells
are streamed (openpyxl's read-only mode), and the layout that mode leaves out is read from the
sheet's XML only as far as the rows shown, so a huge sheet costs its first thousand rows and the
inflating of the rest. (openpyxl itself reads a sheet through once when it lacks the `<dimension>`
Excel writes at its top.)
"""

import colorsys
import io
import logging
import math
import re
import warnings
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from typing import IO, Final, Literal
from xml.etree.ElementTree import Element, XMLPullParser

import openpyxl
import xlrd
from openpyxl.cell.read_only import ReadOnlyCell
from openpyxl.styles.colors import COLOR_INDEX, Color
from openpyxl.styles.fills import PatternFill
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.xml.functions import fromstring

from app.platform.filetypes import XLS, XLSX

logger = logging.getLogger(__name__)
# openpyxl warns about every styling feature it skips (a header, a data validation); none of them
# changes a cell's value, and a warning per file would bury the log.
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

MAX_SHEETS: Final = 20
MAX_ROWS: Final = 1000
MAX_COLUMNS: Final = 50
# An .xlsx is zipped XML: one that inflates past this is refused before it is parsed (a zip bomb).
_MAX_UNZIPPED: Final = 200 * 1024 * 1024
_CHUNK: Final = 64 * 1024
_DEFAULT_FONT_SIZE: Final = 11.0
# Excel's defaults at its default font, Calibri 11, whose digits are 7 pixels wide: a column of 64
# pixels and a row of 15 points.
_DIGIT: Final = 7
_DEFAULT_WIDTH: Final = 64
_DEFAULT_HEIGHT: Final = 15.0
# A font name the web app can ask the reader's computer for; anything else is left out.
_FONT_NAME: Final = re.compile(r"[\w .-]{1,64}")

CellKind = Literal["number", "date", "text", "bool", "error", "blank"]
Border = Literal["thin", "medium", "thick", "double", "dashed", "dotted"]
HAlign = Literal["left", "center", "right", "justify"]
VAlign = Literal["top", "middle", "bottom"]


@dataclass(frozen=True, slots=True)
class Style:
    bold: bool = False
    italic: bool = False
    underline: bool = False
    strike: bool = False
    # None is the workbook's default font and size.
    font: str | None = None
    # Points.
    size: float | None = None
    # "#rrggbb"; None is the default: black text, no fill.
    color: str | None = None
    fill: str | None = None
    # None is Excel's General: numbers to the right, text to the left.
    align: HAlign | None = None
    # None is Excel's default, the bottom.
    valign: VAlign | None = None
    wrap: bool = False
    indent: int = 0
    top: Border | None = None
    right: Border | None = None
    bottom: Border | None = None
    left: Border | None = None


@dataclass(frozen=True, slots=True)
class Cell:
    kind: CellKind
    # A number as a float; a date as ISO text (`2026-09-28T14:30:00`); text, TRUE/FALSE or an error
    # (`#DIV/0!`) as text; "" for a blank cell that has only a style (a border, a fill).
    value: str | float
    # Excel's number format code: `General`, `#,##0.00`, `0%`, `dd.mm.yyyy`.
    format: str
    # Index into the sheet's styles; 0 is the default style.
    style: int


@dataclass(frozen=True, slots=True)
class Row:
    cells: list[Cell | None]
    # Points; None is the sheet's default.
    height: float | None = None
    hidden: bool = False


@dataclass(frozen=True, slots=True)
class Column:
    # Pixels at 100%; None is the sheet's default.
    width: int | None = None
    hidden: bool = False


@dataclass(frozen=True, slots=True)
class Merge:
    # Zero-based, from the top-left cell.
    row: int
    column: int
    rows: int
    columns: int


@dataclass(frozen=True, slots=True)
class Sheet:
    name: str
    columns: list[Column]
    rows: list[Row]
    merges: list[Merge]
    styles: list[Style]
    # The sheet has more rows or columns than were read.
    truncated: bool
    # A column's width in pixels and a row's height in points, where the sheet sets none of its own.
    default_width: int = _DEFAULT_WIDTH
    default_height: float = _DEFAULT_HEIGHT
    # Hidden by its workbook: Excel shows no tab for it.
    hidden: bool = False


@dataclass(frozen=True, slots=True)
class Workbook:
    sheets: list[Sheet]
    # Sheets past the cap, not read.
    omitted: int
    # The font and size of every cell whose style names none: Excel's Normal style's.
    font: str | None = None
    font_size: float = _DEFAULT_FONT_SIZE


class Unreadable(ValueError):
    pass


class TooLarge(Unreadable):
    pass


def read_workbook(data: bytes, media: str) -> Workbook:
    try:
        if media == XLSX:
            return _xlsx(data)
        if media == XLS:
            return _xls(data)
    except Unreadable:
        raise
    except Exception as error:  # a broken or hostile file fails in many ways; all mean "unreadable"
        logger.info("a spreadsheet could not be read: %s", type(error).__name__)
        raise Unreadable("The spreadsheet cannot be read") from error
    raise Unreadable("Not a spreadsheet")


@dataclass(slots=True)
class _Styles:
    """Each distinct style once, the default first."""

    index: dict[Style, int] = field(default_factory=lambda: {Style(): 0})

    def of(self, style: Style) -> int:
        return self.index.setdefault(style, len(self.index))

    def listed(self) -> list[Style]:
        return list(self.index)


# .xlsx


def _xlsx(data: bytes) -> Workbook:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > _MAX_UNZIPPED:
            raise TooLarge("The spreadsheet is too large to show")
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            theme = _theme_colours(book.loaded_theme)
            normal = _xlsx_normal(book)
            sheets = [
                _xlsx_sheet(archive, sheet, theme, normal)
                for sheet in book.worksheets[:MAX_SHEETS]
                if isinstance(sheet, ReadOnlyWorksheet)
            ]
            return Workbook(
                sheets=sheets,
                omitted=max(0, len(book.worksheets) - MAX_SHEETS),
                font=normal.font,
                font_size=normal.size,
            )
        finally:
            book.close()


@dataclass(frozen=True, slots=True)
class _Normal:
    """The font of Excel's Normal style: every cell without a style of its own is drawn in it, and a
    style names a font or size only where it differs."""

    font: str | None = None
    size: float = _DEFAULT_FONT_SIZE


def _xlsx_normal(book: openpyxl.Workbook) -> _Normal:
    """Cell style 0 is Excel's Normal."""
    # The workbook's cell formats and fonts as openpyxl read them; not in the stubs.
    styles, fonts = book._cell_styles, book._fonts  # type: ignore[attr-defined]
    if not styles or styles[0].fontId >= len(fonts):
        return _Normal()
    font = fonts[styles[0].fontId]
    return _Normal(font=_font(font.name, None), size=float(font.sz or _DEFAULT_FONT_SIZE))


def _xlsx_sheet(
    archive: zipfile.ZipFile, sheet: ReadOnlyWorksheet, theme: list[str], normal: _Normal
) -> Sheet:
    # Read-only mode reads cells alone; the path to the sheet's XML is the one it streams them from.
    layout = _xlsx_layout(archive, sheet._worksheet_path)  # type: ignore[attr-defined]  # not in the stubs
    styles = _Styles()
    rows: list[Row] = []
    truncated = False
    # Rows are padded to the width the sheet declares, which can be Excel's full 16,384 columns.
    wide = (sheet.max_column or 0) > MAX_COLUMNS
    # The stubs type iter_rows for the editable worksheet only; the read-only one has the same.
    for number, cells in enumerate(sheet.iter_rows(), start=1):  # type: ignore[misc]
        if number > MAX_ROWS:
            truncated = True
            break
        if wide and any(cell.value is not None for cell in cells[MAX_COLUMNS:]):
            truncated = True
        height, hidden = layout.rows.get(number, (None, False))
        rows.append(
            Row(
                cells=[_xlsx_cell(cell, styles, theme, normal) for cell in cells[:MAX_COLUMNS]],
                height=height,
                hidden=hidden,
            )
        )
    rows = _trim(rows, layout.merges)
    merges = _clip(layout.merges, len(rows))
    return Sheet(
        name=sheet.title,
        columns=_columns(layout.columns, rows, merges),
        rows=rows,
        merges=merges,
        styles=styles.listed(),
        truncated=truncated or _cut_across(layout.merges, len(rows)),
        default_width=layout.width,
        default_height=layout.height,
        hidden=sheet.sheet_state != "visible",
    )


@dataclass(slots=True)
class _Layout:
    columns: dict[int, Column] = field(default_factory=dict)  # zero-based
    rows: dict[int, tuple[float | None, bool]] = field(default_factory=dict)  # one-based
    merges: list[Merge] = field(default_factory=list)
    width: int = _DEFAULT_WIDTH
    height: float = _DEFAULT_HEIGHT


def _xlsx_layout(archive: zipfile.ZipFile, path: str) -> _Layout:
    """The sheet's default sizes, column widths, row heights and merges. The XML is parsed only as
    far as the rows shown; the merges, listed after every row, are found by a byte search."""
    layout = _Layout()
    parser: XMLPullParser[Element] = XMLPullParser(events=("start", "end"))
    rows: Element | None = None
    number = 0
    shown = False
    with archive.open(path) as source:
        while not shown and (chunk := source.read(_CHUNK)):
            parser.feed(chunk)
            for item in parser.read_events():
                # Only start and end are asked for, and both carry the element.
                event, element = item[0], item[-1]
                if not isinstance(element, Element):
                    continue
                tag = element.tag.rpartition("}")[2]
                if event == "start":
                    if tag == "sheetData":
                        rows = element
                elif tag == "sheetFormatPr":
                    stored, base = element.get("defaultColWidth"), element.get("baseColWidth")
                    layout.width = _default_width(
                        float(stored) if stored else None, float(base) if base else None
                    )
                    height = element.get("defaultRowHeight")
                    layout.height = float(height) if height else _DEFAULT_HEIGHT
                elif tag == "col":
                    width = element.get("width")
                    column = Column(
                        width=_pixels(float(width)) if width else None,
                        hidden=_flag(element, "hidden"),
                    )
                    first, last = int(element.get("min", "1")), int(element.get("max", "1"))
                    # Clamped both ways: a file can say min="-1000000000".
                    for index in range(max(first - 1, 0), min(last, MAX_COLUMNS)):
                        layout.columns[index] = column
                elif tag == "row":
                    # `r` may be left out, meaning the row after the last.
                    number = int(element.get("r") or number + 1)
                    custom = _flag(element, "customHeight")
                    if number <= MAX_ROWS and (custom or _flag(element, "hidden")):
                        height = element.get("ht")
                        layout.rows[number] = (
                            float(height) if height and custom else None,
                            _flag(element, "hidden"),
                        )
                    # A row is dropped from the tree once read, so the memory used stays flat.
                    if rows is not None:
                        rows.clear()
                    shown = shown or number >= MAX_ROWS
                elif tag == "sheetData":
                    shown = True
        # The tag that ended the rows shown closed in this chunk, and the merges follow it.
        if shown:
            layout.merges = _merges_from(chunk, source)
    return layout


_MERGES_START: Final = re.compile(rb"<(?:[\w.-]+:)?mergeCells\b[^>]*?(/?)>")
_MERGES_END: Final = re.compile(rb"</(?:[\w.-]+:)?mergeCells\s*>")
_MERGE_REF: Final = re.compile(rb"<(?:[\w.-]+:)?mergeCell\b[^>]*?\bref\s*=\s*[\"']([^\"']*)[\"']")


def _merges_from(chunk: bytes, source: IO[bytes]) -> list[Merge]:
    """The sheet's `<mergeCells>` list, found in the bytes from `chunk` on without parsing the rows
    before it. Cell text cannot pass for it, as XML writes a `<` in text as `&lt;`; a file faking
    one in a CDATA section fakes only its own preview."""
    text = bytearray(chunk)
    while (start := _MERGES_START.search(text)) is None:
        more = source.read(_CHUNK)
        if not more:
            return []
        # Keep enough of the end to find a start tag cut in two between chunks.
        text = text[-256:] + more
    if start.group(1):  # <mergeCells/>
        return []
    del text[: start.end()]
    searched = 0
    while (end := _MERGES_END.search(text, searched)) is None:
        more = source.read(_CHUNK)
        if not more:
            break
        searched = max(0, len(text) - 64)
        text += more
    listed = text[: end.start()] if end else text
    merges = (_merge(ref.decode("latin-1")) for ref in _MERGE_REF.findall(listed))
    return [merge for merge in merges if merge is not None]


def _pixels(width: float) -> int:
    """A column width as Excel stores it, in digits of the default font plus 5 pixels of padding,
    in pixels (ECMA-376 Part 1, 18.3.1.13)."""
    return int((256 * width + 128 // _DIGIT) / 256 * _DIGIT)


def _default_width(stored: float | None, base: float | None) -> int:
    """A sheet's default column width, in pixels. It is stored as a column's width is; without it,
    Excel pads the base width in digits by 5 pixels and rounds up to 8: 8 digits are 64 pixels."""
    if stored:
        return _pixels(stored)
    if base:
        return math.ceil((base * _DIGIT + 5) / 8) * 8
    return _DEFAULT_WIDTH


def _flag(element: Element, name: str) -> bool:
    return element.get(name) in ("1", "true")


def _merge(ref: str) -> Merge | None:
    """`A1:D4` as a merge; None for anything else, which a file from outside may hold."""
    if not _RANGE.fullmatch(ref):
        return None
    first, _, last = ref.partition(":")
    (top, left), (bottom, right) = _coordinate(first), _coordinate(last)
    if bottom < top or right < left:
        return None
    return Merge(row=top, column=left, rows=bottom - top + 1, columns=right - left + 1)


_RANGE: Final = re.compile(r"[A-Za-z]{1,3}[0-9]{1,7}:[A-Za-z]{1,3}[0-9]{1,7}")


def _coordinate(ref: str) -> tuple[int, int]:
    """`B3` as zero-based (row, column)."""
    letters = ref.rstrip("0123456789")
    column = 0
    for letter in letters.upper():
        column = column * 26 + ord(letter) - 64
    return int(ref[len(letters) :]) - 1, column - 1


def _xlsx_cell(cell: object, styles: _Styles, theme: list[str], normal: _Normal) -> Cell | None:
    if not isinstance(cell, ReadOnlyCell):  # the filler of a missing cell
        return None
    style = styles.of(_xlsx_style(cell, theme, normal)) if cell.has_style else 0
    value = cell.value
    if value is None:
        return Cell(kind="blank", value="", format="General", style=style) if style else None
    number_format = str(cell.number_format or "General")
    match value:
        case bool():
            return Cell(
                kind="bool", value="TRUE" if value else "FALSE", format="General", style=style
            )
        case int() | float():
            return Cell(kind="number", value=float(value), format=number_format, style=style)
        case datetime() | date() | time():
            return Cell(kind="date", value=_iso(value), format=number_format, style=style)
        case timedelta():
            days = value.total_seconds() / 86400
            return Cell(kind="number", value=days, format=number_format, style=style)
        case _:
            kind: CellKind = "error" if cell.data_type == "e" else "text"
            text = str(value).replace("\x00", "")
            return Cell(kind=kind, value=text, format=number_format, style=style)


def _xlsx_style(cell: ReadOnlyCell, theme: list[str], normal: _Normal) -> Style:
    font, fill, alignment, border = cell.font, cell.fill, cell.alignment, cell.border
    size = float(font.sz) if font.sz else None
    fill_colour = None
    if isinstance(fill, PatternFill) and fill.fill_type not in (None, "none"):
        fill_colour = _colour(fill.fgColor, theme)
    return Style(
        bold=bool(font.b),
        italic=bool(font.i),
        underline=font.u not in (None, "none"),
        strike=bool(font.strike),
        font=_font(font.name, normal.font),
        size=None if size in (None, normal.size) else size,
        color=_colour(font.color, theme),
        fill=fill_colour,
        align=_H_ALIGN.get(alignment.horizontal or ""),
        valign=_V_ALIGN.get(alignment.vertical or ""),
        wrap=bool(alignment.wrap_text),
        indent=int(alignment.indent or 0),
        top=_BORDERS.get(border.top.style or "") if border.top else None,
        right=_BORDERS.get(border.right.style or "") if border.right else None,
        bottom=_BORDERS.get(border.bottom.style or "") if border.bottom else None,
        left=_BORDERS.get(border.left.style or "") if border.left else None,
    )


_H_ALIGN: Final[dict[str, HAlign]] = {
    "left": "left",
    "fill": "left",
    "center": "center",
    "centerContinuous": "center",
    "right": "right",
    "justify": "justify",
    "distributed": "justify",
}
_V_ALIGN: Final[dict[str, VAlign]] = {
    "top": "top",
    "center": "middle",
    "justify": "middle",
    "distributed": "middle",
    "bottom": "bottom",
}
_BORDERS: Final[dict[str, Border]] = {
    "hair": "thin",
    "thin": "thin",
    "medium": "medium",
    "thick": "thick",
    "double": "double",
    "dashed": "dashed",
    "mediumDashed": "dashed",
    "dashDot": "dashed",
    "mediumDashDot": "dashed",
    "dashDotDot": "dashed",
    "mediumDashDotDot": "dashed",
    "slantDashDot": "dashed",
    "dotted": "dotted",
}


def _colour(colour: Color | None, theme: list[str]) -> str | None:
    """A colour as `#rrggbb`: given outright, from the workbook's theme, or from Excel's old
    palette. Excel's automatic colour (and the palette's system entries) is the default."""
    if colour is None:
        return None
    if colour.type == "rgb" and isinstance(colour.rgb, str) and len(colour.rgb) >= 6:
        return "#" + colour.rgb[-6:].lower()
    if colour.type == "theme" and isinstance(colour.theme, int) and colour.theme < len(theme):
        return _tinted(theme[colour.theme], float(colour.tint or 0))
    if colour.type == "indexed" and isinstance(colour.indexed, int) and colour.indexed < 64:
        return "#" + COLOR_INDEX[colour.indexed][-6:].lower()
    return None


def _tinted(rgb: str, tint: float) -> str:
    """Excel's tint: towards black below zero, towards white above."""
    if not tint:
        return rgb
    red, green, blue = (int(rgb[index : index + 2], 16) / 255 for index in (1, 3, 5))
    hue, lightness, saturation = colorsys.rgb_to_hls(red, green, blue)
    lightness = lightness * (1 + tint) if tint < 0 else lightness * (1 - tint) + tint
    return "#" + "".join(
        f"{round(part * 255):02x}" for part in colorsys.hls_to_rgb(hue, lightness, saturation)
    )


def _theme_colours(theme: bytes | None) -> list[str]:
    """The theme's colours in the order a cell names them. The theme lists dark 1, light 1, dark 2,
    light 2; a cell's theme index counts light 1, dark 1, light 2, dark 2."""
    if not theme:
        return []
    scheme = next((e for e in fromstring(theme).iter() if e.tag.endswith("}clrScheme")), None)
    if scheme is None:
        return []
    listed = []
    for slot in scheme:
        found = next(iter(slot), None)
        value = None if found is None else found.get("lastClr") or found.get("val")
        listed.append(f"#{value.lower()}" if value and len(value) == 6 else "#000000")
    if len(listed) >= 4:
        listed[0], listed[1], listed[2], listed[3] = listed[1], listed[0], listed[3], listed[2]
    return listed


# .xls


def _xls(data: bytes) -> Workbook:
    book = xlrd.open_workbook(file_contents=data, formatting_info=True, on_demand=True)
    try:
        normal = _xls_normal(book)
        sheets = []
        for index in range(min(book.nsheets, MAX_SHEETS)):
            sheets.append(_xls_sheet(book, book.sheet_by_index(index), normal))
            book.unload_sheet(index)
        return Workbook(
            sheets=sheets,
            omitted=max(0, book.nsheets - MAX_SHEETS),
            font=normal.font,
            font_size=normal.size,
        )
    finally:
        book.release_resources()


def _xls_normal(book: xlrd.Book) -> _Normal:
    """Format record 0 is Excel's Normal."""
    if not book.xf_list or book.xf_list[0].font_index >= len(book.font_list):
        return _Normal()
    font = book.font_list[book.xf_list[0].font_index]
    # Twentieths of a point.
    return _Normal(font=_font(font.name, None), size=font.height / 20 or _DEFAULT_FONT_SIZE)


def _xls_sheet(book: xlrd.Book, sheet: xlrd.sheet.Sheet, normal: _Normal) -> Sheet:
    styles = _Styles()
    xf_styles: dict[int, int] = {}
    rows = []
    for number in range(min(sheet.nrows, MAX_ROWS)):
        info = sheet.rowinfo_map.get(number)
        rows.append(
            Row(
                cells=[
                    _xls_cell(book, cell, styles, xf_styles, normal)
                    for cell in sheet.row(number)[:MAX_COLUMNS]
                ],
                # A height set by hand; otherwise the row fits its font.
                height=info.height / 20 if info is not None and info.height_mismatch else None,
                hidden=bool(info is not None and info.hidden),
            )
        )
    listed = [
        Merge(row=top, column=left, rows=bottom - top, columns=right - left)
        for top, bottom, left, right in sheet.merged_cells
    ]
    rows = _trim(rows, listed)
    merges = _clip(listed, len(rows))
    columns = {
        index: Column(width=_pixels(info.width / 256), hidden=bool(info.hidden))
        for index, info in sheet.colinfo_map.items()
        if index < MAX_COLUMNS
    }
    # STANDARDWIDTH is .xlsx's defaultColWidth in 1/256 of a digit, DEFCOLWIDTH its baseColWidth.
    standard, base = sheet.standardwidth, sheet.defcolwidth
    return Sheet(
        name=sheet.name,
        columns=_columns(columns, rows, merges),
        rows=rows,
        merges=merges,
        styles=styles.listed(),
        truncated=sheet.nrows > MAX_ROWS
        or sheet.ncols > MAX_COLUMNS
        or _cut_across(listed, len(rows)),
        default_width=_default_width(standard / 256 if standard else None, base),
        # Twentieths of a point.
        default_height=sheet.default_row_height / 20
        if sheet.default_row_height
        else _DEFAULT_HEIGHT,
        # 1 is hidden, 2 very hidden (shown by no menu in Excel).
        hidden=sheet.visibility != 0,
    )


def _xls_cell(
    book: xlrd.Book,
    cell: xlrd.sheet.Cell,
    styles: _Styles,
    xf_styles: dict[int, int],
    normal: _Normal,
) -> Cell | None:
    xf_index = cell.xf_index or 0
    if xf_index not in xf_styles:
        xf_styles[xf_index] = styles.of(_xls_style(book, book.xf_list[xf_index], normal))
    style = xf_styles[xf_index]
    number_format = book.format_map[book.xf_list[xf_index].format_key].format_str or "General"
    match cell.ctype:
        case xlrd.XL_CELL_EMPTY | xlrd.XL_CELL_BLANK:
            return Cell(kind="blank", value="", format="General", style=style) if style else None
        case xlrd.XL_CELL_NUMBER:
            return Cell(kind="number", value=float(cell.value), format=number_format, style=style)
        case xlrd.XL_CELL_DATE:
            moment = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
            return Cell(kind="date", value=_iso(moment), format=number_format, style=style)
        case xlrd.XL_CELL_BOOLEAN:
            value = "TRUE" if cell.value else "FALSE"
            return Cell(kind="bool", value=value, format="General", style=style)
        case xlrd.XL_CELL_ERROR:
            text = str(xlrd.error_text_from_code.get(cell.value, "#ERROR"))
            return Cell(kind="error", value=text, format="General", style=style)
        case _:
            text = str(cell.value).replace("\x00", "")
            return Cell(kind="text", value=text, format=number_format, style=style)


def _xls_style(book: xlrd.Book, xf: xlrd.formatting.XF, normal: _Normal) -> Style:
    font = book.font_list[xf.font_index]
    size = font.height / 20
    border = xf.border
    return Style(
        bold=bool(font.bold) or font.weight >= 700,
        italic=bool(font.italic),
        underline=bool(font.underline_type),
        strike=bool(font.struck_out),
        font=_font(font.name, normal.font),
        size=None if size == normal.size else size,
        color=_xls_colour(book, font.colour_index),
        fill=_xls_colour(book, xf.background.pattern_colour_index)
        if xf.background.fill_pattern
        else None,
        align=_XLS_H_ALIGN.get(xf.alignment.hor_align),
        valign=_XLS_V_ALIGN.get(xf.alignment.vert_align),
        wrap=bool(xf.alignment.text_wrapped),
        indent=int(xf.alignment.indent_level),
        top=_XLS_BORDERS.get(border.top_line_style),
        right=_XLS_BORDERS.get(border.right_line_style),
        bottom=_XLS_BORDERS.get(border.bottom_line_style),
        left=_XLS_BORDERS.get(border.left_line_style),
    )


def _xls_colour(book: xlrd.Book, index: int) -> str | None:
    rgb = book.colour_map.get(index)
    return "#{:02x}{:02x}{:02x}".format(*rgb) if rgb else None


_XLS_H_ALIGN: Final[dict[int, HAlign]] = {
    1: "left",
    2: "center",
    3: "right",
    4: "left",
    5: "justify",
    6: "center",
    7: "justify",
}
_XLS_V_ALIGN: Final[dict[int, VAlign]] = {
    0: "top",
    1: "middle",
    2: "bottom",
    3: "middle",
    4: "middle",
}
_XLS_BORDERS: Final[dict[int, Border]] = {
    1: "thin",
    2: "medium",
    3: "dashed",
    4: "dotted",
    5: "thick",
    6: "double",
    7: "thin",
    8: "dashed",
    9: "dashed",
    10: "dashed",
    11: "dashed",
    12: "dashed",
    13: "dashed",
}


# Both formats


def _iso(value: datetime | date | time) -> str:
    """A date as ISO text without a zone, as Excel keeps none; a time alone is on Excel's day 0."""
    match value:
        case datetime():
            return value.replace(tzinfo=None).isoformat(timespec="seconds")
        case date():
            return datetime.combine(value, time()).isoformat(timespec="seconds")
        case time():
            return datetime.combine(date(1899, 12, 30), value).isoformat(timespec="seconds")


def _font(name: object, default: str | None) -> str | None:
    """A font's name; None where it is the workbook's default, or not a plain name."""
    if not isinstance(name, str) or not _FONT_NAME.fullmatch(name) or name == default:
        return None
    return name


def _trim(rows: list[Row], merges: list[Merge]) -> list[Row]:
    """Drops empty cells at the end of each row, and the empty rows at the end of the sheet that no
    merge reaches down into: a title merged over four rows keeps its four rows."""
    reach = max((merge.row + merge.rows for merge in merges if merge.row < len(rows)), default=0)
    trimmed = []
    for row in rows:
        end = len(row.cells)
        while end and row.cells[end - 1] is None:
            end -= 1
        trimmed.append(Row(cells=row.cells[:end], height=row.height, hidden=row.hidden))
    while len(trimmed) > reach and not trimmed[-1].cells:
        trimmed.pop()
    return trimmed


def _columns(known: dict[int, Column], rows: list[Row], merges: list[Merge]) -> list[Column]:
    """One entry per column the cells or merges reach, the sheet's default where it set none."""
    width = max(
        max((len(row.cells) for row in rows), default=0),
        max((merge.column + merge.columns for merge in merges), default=0),
    )
    return [known.get(index, Column()) for index in range(width)]


def _cut_across(merges: list[Merge], rows: int) -> bool:
    """Whether a merge among the rows read runs on past the last column read: a title merged over
    A1:AZ1 is cut even with no value past the cap."""
    return any(
        merge.row < rows and merge.column < MAX_COLUMNS < merge.column + merge.columns
        for merge in merges
    )


def _clip(merges: list[Merge], rows: int) -> list[Merge]:
    """Merges within what was read, cut at its edge."""
    clipped = []
    for merge in merges:
        if merge.row >= rows or merge.column >= MAX_COLUMNS:
            continue
        clipped.append(
            Merge(
                row=merge.row,
                column=merge.column,
                rows=min(merge.rows, rows - merge.row),
                columns=min(merge.columns, MAX_COLUMNS - merge.column),
            )
        )
    return [merge for merge in clipped if merge.rows * merge.columns > 1]


# What a person sees


def visible_text(book: Workbook) -> tuple[str, int]:
    """The cells a person sees in Excel, a sheet at a time and a row a line, cells split by tabs;
    and how many cells with a value were left out because the workbook hides their sheet, row or
    column (ADR-0013). A whole number is written without its ".0", as Excel shows it."""
    lines: list[str] = []
    hidden = 0
    for sheet in book.sheets:
        filled = [
            (row, index, cell)
            for row in sheet.rows
            for index, cell in enumerate(row.cells)
            if cell is not None and cell.value != ""
        ]
        shown = [
            (row, index, cell)
            for row, index, cell in filled
            if not sheet.hidden
            and not row.hidden
            and not (index < len(sheet.columns) and sheet.columns[index].hidden)
        ]
        hidden += len(filled) - len(shown)
        if not shown:
            continue
        lines.append(f"# {sheet.name}")
        current: Row | None = None
        cells: list[str] = []
        for row, _, cell in shown:
            if row is not current:
                if cells:
                    lines.append("\t".join(cells))
                current, cells = row, []
            cells.append(_shown(cell))
        lines.append("\t".join(cells))
    return "\n".join(lines), hidden


def _shown(cell: Cell) -> str:
    if isinstance(cell.value, float) and cell.value.is_integer():
        return str(int(cell.value))
    return str(cell.value).strip()
