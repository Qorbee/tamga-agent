import io
from datetime import datetime

import openpyxl
import pytest

from app.modules.documents import sheets
from app.platform.filetypes import XLS, XLSX


def workbook(*rows: list[object], extra_sheet: bool = False) -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Seferler"
    for row in rows:
        sheet.append(row)
    if extra_sheet:
        book.create_sheet("Boş")
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_cells_come_back_as_the_text_a_person_reads() -> None:
    data = workbook(
        ["Plaka", "Tutar", "Tarih", "Saat", "Ödendi", None],
        ["34 ABC 12", 12.0, datetime(2026, 9, 28), datetime(2026, 9, 28, 14, 30), True],
        ["06 XY 99", 0.1 + 0.2, None, None, False],
        [None, None],
        extra_sheet=True,
    )

    [first, empty] = sheets.read_sheets(data, XLSX)

    assert first.name == "Seferler" and not first.truncated
    assert first.rows == [
        ["Plaka", "Tutar", "Tarih", "Saat", "Ödendi"],
        ["34 ABC 12", "12", "2026-09-28", "2026-09-28 14:30", "TRUE"],
        ["06 XY 99", "0.3", None, None, "FALSE"],
    ]
    assert empty.name == "Boş" and empty.rows == []


def test_a_large_sheet_is_cut_and_says_so(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(sheets, "MAX_ROWS", 2)
    monkeypatch.setattr(sheets, "MAX_COLUMNS", 2)

    [sheet] = sheets.read_sheets(workbook(["a", "b", "c"], ["d"], ["e"]), XLSX)

    assert sheet.truncated and sheet.rows == [["a", "b"], ["d"]]


def test_a_value_far_past_the_last_column_still_marks_the_sheet_cut() -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["A1"] = "Plaka"
    sheet.cell(row=1, column=60, value="note")
    out = io.BytesIO()
    book.save(out)

    [read] = sheets.read_sheets(out.getvalue(), XLSX)

    assert read.truncated and read.rows == [["Plaka"]]


def test_a_broken_file_or_a_zip_bomb_is_unreadable(monkeypatch: pytest.MonkeyPatch) -> None:
    with pytest.raises(sheets.Unreadable):
        sheets.read_sheets(b"PK\x03\x04 not really a zip", XLSX)
    with pytest.raises(sheets.Unreadable):
        sheets.read_sheets(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 not really a workbook", XLS)

    monkeypatch.setattr(sheets, "_MAX_UNZIPPED", 100)
    with pytest.raises(sheets.Unreadable, match="too large"):
        sheets.read_sheets(workbook(["x" * 500]), XLSX)
