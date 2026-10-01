import io

import openpyxl

from app.modules.documents import sheets
from app.modules.documents.reading import PdfText, SheetText
from app.platform import isolated
from app.platform.filetypes import PDF, XLSX
from app.platform.isolated import Outcome, Reason
from app.tests.documents.pdfs import pdf, text

ORDER = "Order no 4500123 for truck 34 ABC 12"


async def test_a_pdf_is_read_without_the_text_its_viewer_hides() -> None:
    page = (
        text(ORDER, y=700)
        + text("IGNOREINVISIBLE", y=680, mode=3)
        + text("IGNORECLIP", y=660, mode=7)
        + text("IGNORETINY", y=640, size=0.5)
        + text("IGNORECLEAR", y=620, transparent=True)
    )

    outcome = await isolated.run("documents.pdf_text", pdf(page, text("Page two", y=700)), PDF)
    read = PdfText.model_validate(outcome.value)

    assert read.page_count == 2
    assert ORDER in read.pages[0] and "Page two" in read.pages[1]
    assert "IGNORE" not in "".join(read.pages)
    assert read.hidden == len("IGNOREINVISIBLE IGNORECLIP IGNORETINY IGNORECLEAR".replace(" ", ""))


async def test_an_outline_counts_by_its_own_transparency() -> None:
    page = (
        text("OUTLINED", y=700, mode=1, transparent=True)
        + text("BOTHWAYS", y=680, mode=2, transparent=True)
        + text("IGNOREOUTLINE", y=660, mode=1, clear_outline=True)
        + text("IGNOREBOTH", y=640, mode=2, transparent=True, clear_outline=True)
    )

    read = PdfText.model_validate((await isolated.run("documents.pdf_text", pdf(page), PDF)).value)

    # A clear fill hides nothing its visible outline draws; a clear outline hides an outline.
    assert "OUTLINED" in read.pages[0] and "BOTHWAYS" in read.pages[0]
    assert "IGNORE" not in read.pages[0]
    assert read.hidden == len("IGNOREOUTLINEIGNOREBOTH")


async def test_a_page_with_no_text_reads_empty_and_a_broken_pdf_is_damaged() -> None:
    scanned = PdfText.model_validate((await isolated.run("documents.pdf_text", pdf(""), PDF)).value)

    assert scanned.pages == [""] and scanned.hidden == 0
    assert await isolated.run("documents.pdf_text", b"%PDF-1.7 a", PDF) == Outcome(
        refused=Reason.DAMAGED
    )


async def test_a_spreadsheet_is_read_without_hidden_sheets_rows_and_columns() -> None:
    book = openpyxl.Workbook()
    shown = book.active
    assert shown is not None
    shown.title = "Seferler"
    shown.append(["Plaka", "Sipariş", "Gizli not"])
    shown.append(["34 ABC 12", 4500123.0, "IGNORE column"])
    shown.append(["IGNORE row", 1.0, "x"])
    shown.column_dimensions["C"].hidden = True
    shown.row_dimensions[3].hidden = True
    secret = book.create_sheet("Gizli")
    secret.append(["IGNORE sheet"])
    secret.sheet_state = "hidden"
    data = io.BytesIO()
    book.save(data)

    outcome = await isolated.run("documents.sheet_text", data.getvalue(), XLSX)
    read = SheetText.model_validate(outcome.value)

    assert read.text == "# Seferler\nPlaka\tSipariş\n34 ABC 12\t4500123"
    # Two cells in the hidden column above the hidden row, three in the row, one on the sheet.
    assert read.hidden == 6 and read.sheet_count == 2 and not read.truncated


async def test_a_spreadsheet_past_what_is_read_says_so() -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.cell(row=sheets.MAX_ROWS + 1, column=1, value="4500123")
    data = io.BytesIO()
    book.save(data)

    outcome = await isolated.run("documents.sheet_text", data.getvalue(), XLSX)

    assert SheetText.model_validate(outcome.value).truncated
